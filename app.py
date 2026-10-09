# Required dependencies for this script to run fully:
# pip install streamlit pandas openpyxl lxml html5lib xlrd
pip install "xlrd>=2.0.1"
import io
import math
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Conciliación Contable Multiorigen",
    page_icon="📊",
    layout="wide"
)

ERP_CONFIGS = {
    "AutoDealer": {
        "asiento": "ASIENTO",
        "fecha": "FECHA",
        "detalle": "DETALLE",
        "referencia": "REFERENCIA",
        "entidad": "ENTIDAD",
        "cuit": "CUIT",
        "debe": "DEBE",
        "haber": "HABER",
        "saldo": "SALDO",
    },
    "Bejerman": {
        "asiento": "AsientoNumero",
        "fecha": "Fecha",
        "detalle": "Concepto",
        "referencia": "CodigoCliProv",
        "ref_secundaria": "NroCompFlex",
        "entidad": "RazonSocial",
        "cuit": "CliProvNroDocumento",
        "debe": "ImpDebe_Loc",
        "haber": "ImpHaber_Loc",
        "saldo": "ImpTotal_Loc",
    }
}

SIN_RAZON_SOCIAL = "NO ENCONTRADO EN PADRÓN"


def safe_float(val):
    """
    Convierte de forma segura cualquier valor (NaN, None, cadenas formateadas) a float.
    Evita los falsos positivos de evaluación booleana de NaN en Python.
    """
    if pd.isna(val) or val is None:
        return 0.0
    try:
        if isinstance(val, str):
            val = val.replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
        v = float(val)
        return 0.0 if math.isnan(v) else v
    except (ValueError, TypeError):
        return 0.0


def clean_cuit(cuit_val):
    """
    Limpia y normaliza un CUIT/CUIL dejando solo dígitos numéricos.
    Mantiene ceros a la izquierda si los hay.
    """
    if pd.isna(cuit_val) or cuit_val is None:
        return ""
    # Convertir a string sin decimales flotantes .0
    s_cuit = str(cuit_val).replace(".0", "").strip()
    digits = re.sub(r'\D', '', s_cuit)
    return digits if len(digits) >= 8 else s_cuit.strip()


def clean_entity_key(nombre):
    """
    Sanitiza y normaliza el nombre de la Razón Social / Entidad como criterio secundario.
    """
    if not nombre or pd.isna(nombre):
        return ""
    s = str(nombre).upper()
    s = re.sub(r'[^A-Z0-9]', ' ', s)
    words = s.split()
    noise_words = {"SA", "SAU", "SRL", "SACIF", "LIMITADA", "LTD", "INC", "CORP", "SOCIEDAD", "ANONIMA"}
    filtered = [w for w in words if w not in noise_words]
    return " ".join(filtered) if filtered else " ".join(words)


def buscar_columna_arca(df_arca, posibles_nombres):
    """
    Busca de forma flexible la columna correspondiente en el reporte de ARCA.
    """
    for col in df_arca.columns:
        col_clean = str(col).strip().lower()
        for p in posibles_nombres:
            if p.lower() in col_clean:
                return col
    return posibles_nombres[0] if posibles_nombres else ""


def obtener_monto_mayor(row, cols_config, priorizar_total=False):
    """
    Obtiene el monto del registro del Mayor evaluando DEBE, HABER y SALDO
    según la configuración de columnas del ERP activo.

    Si priorizar_total=True (módulo de percepciones IIBB) se toma primero la columna
    de total con signo (ImpTotal_Loc en Bejerman), donde las NC ya vienen en negativo.
    Si esa columna está vacía/0 se cae a la lógica habitual DEBE / HABER.
    """
    col_debe = cols_config.get("debe", "DEBE")
    col_haber = cols_config.get("haber", "HABER")
    col_saldo = cols_config.get("saldo", "SALDO")

    debe = safe_float(row.get(col_debe))
    haber = safe_float(row.get(col_haber))
    saldo = safe_float(row.get(col_saldo))

    if priorizar_total and saldo != 0:
        return round(saldo, 2)

    if debe != 0:
        return round(debe, 2)
    elif haber != 0:
        return round(-abs(haber), 2)
    elif saldo != 0:
        return round(saldo, 2)

    return 0.0

def validar_formato_archivo(uploaded_file):
    """
    Verifica que el archivo sea un Excel válido (.xlsx).
    Si es .xls o no se puede leer, lanza un error amigable.
    """
    filename = uploaded_file.name.lower()

    if filename.endswith('.xls'):
        raise ValueError(
            f"El archivo '{uploaded_file.name}' está en formato '.xls' (formato antiguo de Excel o archivo web de AFIP/ARCA). "
            f"Por favor, abre el archivo en Excel y guárdalo usando la opción 'Guardar como...' eligiendo el formato "
            f"'Libro de Excel (*.xlsx)' antes de subirlo aquí."
        )

    # Verificamos si realmente es un xlsx intentando leer una pequeña parte
    try:
        # Intentamos leer solo la primera fila para validar el formato rápido
        pd.read_excel(uploaded_file, nrows=1)
    except Exception as e:
         raise ValueError(
            f"El archivo '{uploaded_file.name}' parece estar corrupto o no es un formato de Excel válido (.xlsx). "
            f"Si lo descargaste de ARCA/AFIP, por favor ábrelo en Excel y guárdalo como 'Libro de Excel (*.xlsx)'."
        )

    # Reseteamos el puntero del archivo después de la prueba de lectura
    uploaded_file.seek(0)
    return uploaded_file


# ---------------------------------------------------------------------------
# MÓDULO NUEVO: PERCEPCIONES IIBB (txt del sistema)
# ---------------------------------------------------------------------------

# Formato por defecto del sistema (ancho fijo):
# JJJ CUIT(13) FECHA(10) PV(4) NRO(8) COMPROBANTE(1) TIPO(1) MONTO
# Ej: 90130-50000173-501/10/2025700000095099O 00023732,45
RE_PERCEPCION_FIJO = re.compile(
    r'^(\d{3})(\d{2}-\d{8}-\d)(\d{2}/\d{2}/\d{4})(\d{4})(\d{8})([A-Z])([A-Z ])(\S+)$'
)
RE_CUIT_GUIONES = re.compile(r'^\d{2}-\d{8}-\d$')


def decodificar_texto(raw):
    """Decodifica bytes de un txt: UTF-8 y, si falla, Latin-1 (típico de Bejerman/ERPs)."""
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="ignore")


def parse_monto_txt(s):
    """
    Convierte '00023732,45' / '4755,6' / '-0001234,50' a float.
    Las notas de crédito vienen con signo negativo: se respeta el signo.
    """
    s = str(s).strip()
    negativo = "-" in s or "(" in s
    t = re.sub(r'[^\d,.]', '', s)
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return 0.0
    return -abs(v) if negativo else v


def _fecha_a_dt(fecha_str):
    return datetime.strptime(fecha_str.strip(), "%d/%m/%Y")


def parse_percepciones_txt(raw):
    """
    Interpreta el .txt de percepciones IIBB y devuelve un DataFrame con:
    Jurisdicción, CUIT, Fecha, Nro Referencia, Nro Comprobante, Comprobante,
    Tipo Comprobante, Monto.

    Lee el formato de ancho fijo (primer bloque del sistema). Si el archivo no tiene
    líneas en ese formato, intenta leer el formato ya tabulado (separado por tabs).
    Si el archivo trae ambos bloques, se usa solo el de ancho fijo para no duplicar.
    """
    lineas = decodificar_texto(raw).splitlines()
    registros = []

    # 1) Formato de ancho fijo
    for linea in lineas:
        m = RE_PERCEPCION_FIJO.match(linea.strip())
        if not m:
            continue
        jur, cuit, fecha, ref, nro, comp, tipo, monto = m.groups()
        registros.append((jur, cuit, fecha, int(ref), int(nro), comp, tipo.strip(), parse_monto_txt(monto)))

    # 2) Fallback: formato tabulado
    if not registros:
        for linea in lineas:
            partes = linea.split("\t")
            if len(partes) < 8:
                continue
            jur, cuit = partes[0].strip(), partes[1].strip()
            if not (re.fullmatch(r'\d{3}', jur) and RE_CUIT_GUIONES.match(cuit)):
                continue
            try:
                fecha = _fecha_a_dt(partes[2]).strftime("%d/%m/%Y")
                registros.append((jur, cuit, fecha, int(partes[3]), int(partes[4]),
                                  partes[5].strip(), partes[6].strip(), parse_monto_txt(partes[7])))
            except ValueError:
                continue

    df = pd.DataFrame(registros, columns=[
        "Jurisdicción", "CUIT", "Fecha", "Nro Referencia", "Nro Comprobante",
        "Comprobante", "Tipo Comprobante", "Monto",
    ])
    return df


def cargar_padron_proveedores(raw):
    """
    Lee el listado de proveedores (txt separado por tabs: Cód., Razón Social,
    Tipo Docum., Nro. Docum.) y devuelve {cuit_solo_digitos: razon_social}.
    Si un CUIT está repetido se usa el primero.
    """
    texto = decodificar_texto(raw)
    df = pd.read_csv(io.StringIO(texto), sep="\t", dtype=str)
    col_rs = next((c for c in df.columns if "raz" in str(c).lower()), df.columns[1])
    col_doc = next((c for c in df.columns if "nro" in str(c).lower()), df.columns[-1])

    padron = {}
    for rs, doc in zip(df[col_rs], df[col_doc]):
        k = clean_cuit(doc)
        if k and k not in padron and pd.notna(rs):
            padron[k] = str(rs).strip()
    return padron


def agregar_razon_social(df, padron):
    df = df.copy()
    df["Razón Social"] = df["CUIT"].apply(lambda c: padron.get(clean_cuit(c), SIN_RAZON_SOCIAL))
    return df


def construir_df_motor(df):
    """
    Adapta el DataFrame de percepciones al formato que consume el motor de conciliación
    (mismas columnas que un reporte de ARCA), para que la lógica de cruce sea idéntica.
    El orden de columnas importa: 'Número Comprobante' va antes que 'Descripción Comprobante'.
    """
    df = df.reset_index(drop=True)
    return pd.DataFrame({
        "CUIT Agente Ret./Perc.": df["CUIT"],
        "Número Certificado": df["Nro Comprobante"].astype(str),
        "Fecha Ret./Perc.": df["Fecha"].apply(lambda f: _fecha_a_dt(f).strftime("%Y-%m-%d")),
        "Número Comprobante": (
            df["Nro Referencia"].astype(str).str.zfill(4) + "-" + df["Nro Comprobante"].astype(str).str.zfill(8)
        ),
        "Descripción Comprobante": df["Comprobante"].astype(str) + df["Tipo Comprobante"].astype(str),
        "Denominación o Razón Social": df["Razón Social"],
        "Importe Ret./Perc.": df["Monto"],
    })


# ---------------------------------------------------------------------------
# MOTOR DE CONCILIACIÓN
# ---------------------------------------------------------------------------

def procesar_archivos(
    file_m,
    file_a,
    erp_seleccionado,
    df_arca_override=None,
    origen_label="ARCA",
    concepto="RETENCIONES",
    subtitulo="",
    priorizar_total_mayor=False,
):
    """
    Concilia el Mayor contra un reporte externo.

    - Modo ARCA (por defecto): file_a es el Excel de MisRetenciones.
    - Modo IIBB: se pasa df_arca_override (ya normalizado con construir_df_motor),
      origen_label="IIBB", concepto="PERCEPCIONES IIBB" y priorizar_total_mayor=True.
    La lógica de cruce (1 a 1, lote por CUIT, lote por Razón Social) es la misma en ambos.
    """
    config_erp = ERP_CONFIGS.get(erp_seleccionado, ERP_CONFIGS["AutoDealer"])
    ORI = origen_label

    # Validamos los archivos Excel ANTES de procesarlos
    validar_formato_archivo(file_m)
    df_mayor = pd.read_excel(file_m).dropna(how="all")

    if df_arca_override is None:
        validar_formato_archivo(file_a)
        df_arca = pd.read_excel(file_a).dropna(how="all")

        # Mapeo dinámico de nombres de columnas de ARCA
        col_cuit_arca = buscar_columna_arca(df_arca, ["CUIT Agente Ret./Perc.", "CUIT Agente", "CUIT"])
        col_cert_arca = buscar_columna_arca(df_arca, ["Número Certificado", "Nro Certificado", "Certificado"])
        col_fecha_arca = buscar_columna_arca(df_arca, ["Fecha Ret./Perc.", "Fecha Registro", "Fecha"])
        col_comp_arca = buscar_columna_arca(df_arca, ["Número Comprobante", "Nro Comprobante", "Comprobante"])
        col_razon_arca = buscar_columna_arca(df_arca, ["Denominación o Razón Social", "Razon Social", "Denominacion"])
        col_monto_arca = buscar_columna_arca(df_arca, ["Importe Ret./Perc.", "Importe", "Monto"])
        col_tipo_arca = buscar_columna_arca(df_arca, ["Descripción Comprobante", "Tipo Comprobante"])
    else:
        df_arca = df_arca_override.copy().dropna(how="all")
        col_cuit_arca = "CUIT Agente Ret./Perc."
        col_cert_arca = "Número Certificado"
        col_fecha_arca = "Fecha Ret./Perc."
        col_comp_arca = "Número Comprobante"
        col_razon_arca = "Denominación o Razón Social"
        col_monto_arca = "Importe Ret./Perc."
        col_tipo_arca = "Descripción Comprobante"

    # Columnas del Mayor según el ERP
    col_asiento_m = config_erp.get("asiento", "ASIENTO")
    col_fecha_m = config_erp.get("fecha", "FECHA")
    col_detalle_m = config_erp.get("detalle", "DETALLE")
    col_ref_m = config_erp.get("referencia", "REFERENCIA")
    col_ref_sec_m = config_erp.get("ref_secundaria", "")
    col_entidad_m = config_erp.get("entidad", "ENTIDAD")
    col_cuit_m = config_erp.get("cuit", "CUIT")

    # Calculamos montos numéricos limpios
    df_mayor["MONTO_CALC"] = df_mayor.apply(
        lambda r: obtener_monto_mayor(r, config_erp, priorizar_total_mayor), axis=1
    )
    df_arca["MONTO_CALC"] = df_arca.apply(lambda r: round(safe_float(r.get(col_monto_arca)), 2), axis=1)

    # Filtrar únicamente registros no nulos en monto
    df_mayor = df_mayor[df_mayor["MONTO_CALC"] != 0].copy()
    df_arca = df_arca[df_arca["MONTO_CALC"] != 0].copy()

    # Normalización de CUIT y claves
    df_mayor["CUIT_CLEAN"] = df_mayor.get(col_cuit_m, pd.Series(dtype=object)).apply(clean_cuit)
    df_arca["CUIT_CLEAN"] = df_arca.get(col_cuit_arca, pd.Series(dtype=object)).apply(clean_cuit)

    df_mayor["ENTITY_KEY"] = df_mayor.get(col_entidad_m, pd.Series(dtype=object)).apply(clean_entity_key)
    df_arca["ENTITY_KEY"] = df_arca.get(col_razon_arca, pd.Series(dtype=object)).apply(clean_entity_key)

    df_mayor["ASIENTO_STR"] = (
        df_mayor.get(col_asiento_m, pd.Series(dtype=object))
        .fillna("")
        .astype(str)
        .str.replace(".0", "", regex=False)
    )
    df_arca["CERT_STR"] = (
        df_arca.get(col_cert_arca, pd.Series(dtype=object))
        .fillna("")
        .astype(str)
        .str.replace(".0", "", regex=False)
    )

    conciliadas_1a1 = []
    conciliadas_lote = []
    pendientes_arca = []
    pendientes_mayor = []

    arca_matched_indices = set()
    mayor_matched_indices = set()

    for idx_a, row_a in df_arca.iterrows():
        monto_a = row_a["MONTO_CALC"]
        cuit_a = row_a["CUIT_CLEAN"]

        for idx_m, row_m in df_mayor.iterrows():
            if idx_m in mayor_matched_indices:
                continue

            monto_m = row_m["MONTO_CALC"]
            cuit_m = row_m["CUIT_CLEAN"]

            # Comprobación de coincidencia por monto (y CUIT si ambos están disponibles)
            monto_coincide = abs(monto_a - monto_m) < 0.01
            cuit_coincide = (cuit_a == cuit_m) if (cuit_a and cuit_m) else True

            if monto_coincide and cuit_coincide:
                arca_matched_indices.add(idx_a)
                mayor_matched_indices.add(idx_m)

                # Construir detalle del comprobante del mayor
                detalle_comp = str(row_m.get(col_detalle_m, ""))
                if col_ref_sec_m and row_m.get(col_ref_sec_m):
                    detalle_comp += f" [{row_m.get(col_ref_sec_m)}]"

                conciliadas_1a1.append({
                    "cert_arca": row_a["CERT_STR"],
                    "fecha_arca": str(row_a.get(col_fecha_arca, ""))[:10],
                    "comp_arca": str(row_a.get(col_comp_arca, "")),
                    "cuit_arca": row_a.get(col_cuit_arca, ""),
                    "razon_arca": row_a.get(col_razon_arca, ""),
                    "monto_arca": monto_a,
                    "asiento_mayor": row_m["ASIENTO_STR"],
                    "fecha_mayor": str(row_m.get(col_fecha_m, ""))[:10],
                    "comp_mayor": detalle_comp,
                    "cuit_mayor": row_m.get(col_cuit_m, ""),
                    "razon_mayor": row_m.get(col_entidad_m, ""),
                    "monto_mayor": monto_m,
                })
                break

    unmatched_arca = df_arca[~df_arca.index.isin(arca_matched_indices)].copy()
    unmatched_mayor = df_mayor[~df_mayor.index.isin(mayor_matched_indices)].copy()

    # Prioridad 1 para cruce en lote: CUITs comunes que no estén vacíos
    cuits_arca = set(unmatched_arca["CUIT_CLEAN"]) - {""}
    cuits_mayor = set(unmatched_mayor["CUIT_CLEAN"]) - {""}
    common_cuits = cuits_arca.intersection(cuits_mayor)

    for cuit_key in common_cuits:
        group_arca = unmatched_arca[unmatched_arca["CUIT_CLEAN"] == cuit_key]
        group_mayor = unmatched_mayor[unmatched_mayor["CUIT_CLEAN"] == cuit_key]

        sum_arca = round(group_arca["MONTO_CALC"].sum(), 2)
        sum_mayor = round(group_mayor["MONTO_CALC"].sum(), 2)

        if abs(sum_arca - sum_mayor) < 0.50 and sum_arca != 0:
            for idx in group_arca.index:
                arca_matched_indices.add(idx)
            for idx in group_mayor.index:
                mayor_matched_indices.add(idx)

            certs_list = sorted(list(set(group_arca["CERT_STR"].astype(str))))
            certs_str = ", ".join(certs_list)

            asientos_list = sorted(list(set(group_mayor["ASIENTO_STR"].astype(str))))
            asientos_str = ", ".join(asientos_list)

            fechas_a = sorted([str(f)[:10] for f in group_arca[col_fecha_arca].dropna()])
            fecha_a_str = f"{fechas_a[0]} a {fechas_a[-1]}" if len(fechas_a) > 1 else (fechas_a[0] if fechas_a else "")

            fechas_m = sorted([str(f)[:10] for f in group_mayor[col_fecha_m].dropna()])
            fecha_m_str = f"{fechas_m[0]} a {fechas_m[-1]}" if len(fechas_m) > 1 else (fechas_m[0] if fechas_m else "")

            nombre_entidad = group_arca.iloc[0].get(col_razon_arca) or group_mayor.iloc[0].get(col_entidad_m)
            cuit_real = group_arca.iloc[0].get(col_cuit_arca) or group_mayor.iloc[0].get(col_cuit_m) or cuit_key

            conciliadas_lote.append({
                "empresa": nombre_entidad,
                "cuit": cuit_real,
                "certs_arca": certs_str,
                "fecha_arca": fecha_a_str,
                "monto_arca": sum_arca,
                "asientos_mayor": asientos_str,
                "fecha_mayor": fecha_m_str,
                "monto_mayor": sum_mayor,
                "obs": f"Conciliación por CUIT por lote ({len(group_arca)} reg. {ORI} vs {len(group_mayor)} reg. Mayor)",
            })

    # Prioridad 2 para cruce en lote: Razón social sanitizada para los no matcheados por CUIT
    unmatched_arca_p2 = df_arca[~df_arca.index.isin(arca_matched_indices)].copy()
    unmatched_mayor_p2 = df_mayor[~df_mayor.index.isin(mayor_matched_indices)].copy()

    common_entities = set(unmatched_arca_p2["ENTITY_KEY"]).intersection(set(unmatched_mayor_p2["ENTITY_KEY"])) - {""}

    for ent_key in common_entities:
        group_arca = unmatched_arca_p2[unmatched_arca_p2["ENTITY_KEY"] == ent_key]
        group_mayor = unmatched_mayor_p2[unmatched_mayor_p2["ENTITY_KEY"] == ent_key]

        sum_arca = round(group_arca["MONTO_CALC"].sum(), 2)
        sum_mayor = round(group_mayor["MONTO_CALC"].sum(), 2)

        if abs(sum_arca - sum_mayor) < 0.50 and sum_arca != 0:
            for idx in group_arca.index:
                arca_matched_indices.add(idx)
            for idx in group_mayor.index:
                mayor_matched_indices.add(idx)

            certs_list = sorted(list(set(group_arca["CERT_STR"].astype(str))))
            certs_str = ", ".join(certs_list)

            asientos_list = sorted(list(set(group_mayor["ASIENTO_STR"].astype(str))))
            asientos_str = ", ".join(asientos_list)

            fechas_a = sorted([str(f)[:10] for f in group_arca[col_fecha_arca].dropna()])
            fecha_a_str = f"{fechas_a[0]} a {fechas_a[-1]}" if len(fechas_a) > 1 else (fechas_a[0] if fechas_a else "")

            fechas_m = sorted([str(f)[:10] for f in group_mayor[col_fecha_m].dropna()])
            fecha_m_str = f"{fechas_m[0]} a {fechas_m[-1]}" if len(fechas_m) > 1 else (fechas_m[0] if fechas_m else "")

            nombre_entidad = group_arca.iloc[0].get(col_razon_arca) or group_mayor.iloc[0].get(col_entidad_m)
            cuit_real = group_arca.iloc[0].get(col_cuit_arca) or group_mayor.iloc[0].get(col_cuit_m) or ""

            conciliadas_lote.append({
                "empresa": nombre_entidad,
                "cuit": cuit_real,
                "certs_arca": certs_str,
                "fecha_arca": fecha_a_str,
                "monto_arca": sum_arca,
                "asientos_mayor": asientos_str,
                "fecha_mayor": fecha_m_str,
                "monto_mayor": sum_mayor,
                "obs": f"Conciliación por Razón Social por lote ({len(group_arca)} reg. {ORI} vs {len(group_mayor)} reg. Mayor)",
            })

    final_unmatched_arca = df_arca[~df_arca.index.isin(arca_matched_indices)]
    final_unmatched_mayor = df_mayor[~df_mayor.index.isin(mayor_matched_indices)]

    for idx_a, row_a in final_unmatched_arca.iterrows():
        pendientes_arca.append({
            "cert_arca": row_a["CERT_STR"],
            "fecha_arca": str(row_a.get(col_fecha_arca, ""))[:10],
            "tipo_comp": str(row_a.get(col_tipo_arca, "")),
            "comp_arca": str(row_a.get(col_comp_arca, "")),
            "cuit_arca": str(row_a.get(col_cuit_arca, "")),
            "razon_arca": row_a.get(col_razon_arca, ""),
            "monto_arca": row_a["MONTO_CALC"],
        })

    for idx_m, row_m in final_unmatched_mayor.iterrows():
        ref_val = str(row_m.get(col_ref_m, ""))
        if col_ref_sec_m and row_m.get(col_ref_sec_m):
            ref_val += f" / {row_m.get(col_ref_sec_m)}"

        pendientes_mayor.append({
            "asiento_mayor": row_m["ASIENTO_STR"],
            "fecha_mayor": str(row_m.get(col_fecha_m, ""))[:10],
            "detalle_mayor": str(row_m.get(col_detalle_m, "")),
            "ref_mayor": ref_val,
            "cuit_mayor": str(row_m.get(col_cuit_m, "")),
            "razon_mayor": row_m.get(col_entidad_m, ""),
            "monto_mayor": row_m["MONTO_CALC"],
        })

    sheet_pend_arca = f"Pendientes_{ORI}"
    sheet_pend_mayor = "Pendientes_MAYOR"

    wb = openpyxl.Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen General"
    ws_1a1 = wb.create_sheet(title="Conciliadas 1a1")
    ws_lote = wb.create_sheet(title="Conciliadas por Lote")
    ws_arca = wb.create_sheet(title=sheet_pend_arca)
    ws_mayor = wb.create_sheet(title=sheet_pend_mayor)

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=14, bold=True, color="1F4E79")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)

    fill_resumen_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_1a1_hdr = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_lote_hdr = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_arca_hdr = PatternFill(start_color="C65911", end_color="C65911", fill_type="solid")
    fill_mayor_hdr = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
    fill_total = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    double_bottom = Side(border_style="double", color="000000")
    top_thin = Side(border_style="thin", color="000000")
    total_border = Border(top=top_thin, bottom=double_bottom)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    fmt_currency = '"$ "#,##0.00;("$ "#,##0.00);"-"'

    sufijo_titulo = f" - {subtitulo}" if subtitulo else ""
    lbl_cert = "Certif." if ORI == "ARCA" else "Nro Comp."
    lbl_cert_largo = "Nro Certificado" if ORI == "ARCA" else "Nro Comprobante"

    ws_1a1.cell(row=1, column=1, value=f"PARTIDAS CONCILIADAS EXACTAS 1 A 1 ({erp_seleccionado.upper()}){sufijo_titulo}").font = font_title
    headers_1a1 = [
        f"{lbl_cert} {ORI}",
        f"Fecha {ORI}",
        f"Comprobante {ORI}",
        f"CUIT {ORI}",
        f"Razón Social {ORI}",
        f"Monto {ORI} ($)",
        "Asiento Mayor",
        "Fecha Mayor",
        "Detalle / Comp. Mayor",
        "CUIT Mayor",
        "Razón Social Mayor",
        "Monto Mayor ($)",
        "Diferencia ($)",
    ]
    for col_idx, text in enumerate(headers_1a1, 1):
        c = ws_1a1.cell(row=3, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_header, fill_1a1_hdr, align_center

    r_idx = 4
    for item in conciliadas_1a1:
        ws_1a1.cell(row=r_idx, column=1, value=item["cert_arca"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=2, value=item["fecha_arca"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=3, value=item["comp_arca"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=4, value=item["cuit_arca"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=5, value=item["razon_arca"]).alignment = align_left

        c_m1 = ws_1a1.cell(row=r_idx, column=6, value=item["monto_arca"])
        c_m1.number_format, c_m1.alignment = fmt_currency, align_right

        ws_1a1.cell(row=r_idx, column=7, value=item["asiento_mayor"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=8, value=item["fecha_mayor"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=9, value=item["comp_mayor"]).alignment = align_left
        ws_1a1.cell(row=r_idx, column=10, value=item["cuit_mayor"]).alignment = align_center
        ws_1a1.cell(row=r_idx, column=11, value=item["razon_mayor"]).alignment = align_left

        c_m2 = ws_1a1.cell(row=r_idx, column=12, value=item["monto_mayor"])
        c_m2.number_format, c_m2.alignment = fmt_currency, align_right

        c_diff = ws_1a1.cell(row=r_idx, column=13, value=f"=F{r_idx}-L{r_idx}")
        c_diff.number_format, c_diff.alignment = fmt_currency, align_right

        for c in range(1, 14):
            ws_1a1.cell(row=r_idx, column=c).font = font_regular
            ws_1a1.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

    ws_1a1.cell(row=r_idx, column=1, value="TOTAL CONCILIADO 1 A 1").font = font_bold
    ws_1a1.cell(row=r_idx, column=1).alignment = align_left
    c_tot1 = ws_1a1.cell(row=r_idx, column=6, value=f"=SUM(F4:F{r_idx-1})")
    c_tot1.font, c_tot1.number_format = font_bold, fmt_currency
    c_tot2 = ws_1a1.cell(row=r_idx, column=12, value=f"=SUM(L4:L{r_idx-1})")
    c_tot2.font, c_tot2.number_format = font_bold, fmt_currency
    c_totdiff = ws_1a1.cell(row=r_idx, column=13, value=f"=SUM(M4:M{r_idx-1})")
    c_totdiff.font, c_totdiff.number_format = font_bold, fmt_currency

    for c in range(1, 14):
        ws_1a1.cell(row=r_idx, column=c).fill = fill_total
        ws_1a1.cell(row=r_idx, column=c).border = total_border

    row_1a1_total = r_idx
    cnt_1a1 = len(conciliadas_1a1)

    ws_lote.cell(row=1, column=1, value=f"PARTIDAS CONCILIADAS POR LOTE / CUIT ({erp_seleccionado.upper()}){sufijo_titulo}").font = font_title
    headers_lote = [
        "Razón Social / Empresa",
        "CUIT",
        f"{'Certificados' if ORI == 'ARCA' else 'Comprobantes'} {ORI} Incluidos",
        f"Rango Fecha {ORI}",
        f"Monto Total {ORI} ($)",
        "Asientos Mayor Incluidos",
        "Rango Fecha Mayor",
        "Monto Total Mayor ($)",
        "Diferencia ($)",
        "Observaciones",
    ]
    for col_idx, text in enumerate(headers_lote, 1):
        c = ws_lote.cell(row=4, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_header, fill_lote_hdr, align_center

    r_lote_idx = 5
    for item in conciliadas_lote:
        ws_lote.cell(row=r_lote_idx, column=1, value=item["empresa"]).alignment = align_left
        ws_lote.cell(row=r_lote_idx, column=2, value=item["cuit"]).alignment = align_center
        ws_lote.cell(row=r_lote_idx, column=3, value=item["certs_arca"]).alignment = align_center
        ws_lote.cell(row=r_lote_idx, column=4, value=item["fecha_arca"]).alignment = align_center

        c_m1 = ws_lote.cell(row=r_lote_idx, column=5, value=item["monto_arca"])
        c_m1.number_format, c_m1.alignment = fmt_currency, align_right

        ws_lote.cell(row=r_lote_idx, column=6, value=item["asientos_mayor"]).alignment = align_center
        ws_lote.cell(row=r_lote_idx, column=7, value=item["fecha_mayor"]).alignment = align_center

        c_m2 = ws_lote.cell(row=r_lote_idx, column=8, value=item["monto_mayor"])
        c_m2.number_format, c_m2.alignment = fmt_currency, align_right

        c_diff = ws_lote.cell(row=r_lote_idx, column=9, value=f"=E{r_lote_idx}-H{r_lote_idx}")
        c_diff.number_format, c_diff.alignment = fmt_currency, align_right

        ws_lote.cell(row=r_lote_idx, column=10, value=item["obs"]).alignment = align_left

        for c in range(1, 11):
            ws_lote.cell(row=r_lote_idx, column=c).font = font_regular
            ws_lote.cell(row=r_lote_idx, column=c).border = thin_border
        r_lote_idx += 1

    ws_lote.cell(row=r_lote_idx, column=1, value="TOTAL CONCILIADO POR LOTE").font = font_bold

    if r_lote_idx > 5:
        c_lt1 = ws_lote.cell(row=r_lote_idx, column=5, value=f"=SUM(E5:E{r_lote_idx-1})")
        c_lt2 = ws_lote.cell(row=r_lote_idx, column=8, value=f"=SUM(H5:H{r_lote_idx-1})")
        c_ltdiff = ws_lote.cell(row=r_lote_idx, column=9, value=f"=SUM(I5:I{r_lote_idx-1})")
    else:
        c_lt1 = ws_lote.cell(row=r_lote_idx, column=5, value=0)
        c_lt2 = ws_lote.cell(row=r_lote_idx, column=8, value=0)
        c_ltdiff = ws_lote.cell(row=r_lote_idx, column=9, value=0)

    c_lt1.font, c_lt1.number_format, c_lt1.alignment = font_bold, fmt_currency, align_right
    c_lt2.font, c_lt2.number_format, c_lt2.alignment = font_bold, fmt_currency, align_right
    c_ltdiff.font, c_ltdiff.number_format, c_ltdiff.alignment = font_bold, fmt_currency, align_right

    for c in range(1, 11):
        ws_lote.cell(row=r_lote_idx, column=c).fill = fill_total
        ws_lote.cell(row=r_lote_idx, column=c).border = total_border

    row_lote_total = r_lote_idx
    cnt_lote = len(conciliadas_lote)

    ws_arca.cell(row=1, column=1, value=f"{concepto} PENDIENTES EN {ORI} (NO REGISTRADAS EN MAYOR){sufijo_titulo}").font = font_title
    headers_arca = [
        lbl_cert_largo,
        "Fecha Reg.",
        "Tipo Comprobante",
        "Nro Comprobante",
        "CUIT Agente",
        "Razón Social Agente",
        "Monto ($)" if ORI != "ARCA" else "Monto Retención ($)",
    ]
    for col_idx, text in enumerate(headers_arca, 1):
        c = ws_arca.cell(row=3, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_header, fill_arca_hdr, align_center

    r_arca_idx = 4
    for item in pendientes_arca:
        ws_arca.cell(row=r_arca_idx, column=1, value=item["cert_arca"]).alignment = align_center
        ws_arca.cell(row=r_arca_idx, column=2, value=item["fecha_arca"]).alignment = align_center
        ws_arca.cell(row=r_arca_idx, column=3, value=item["tipo_comp"]).alignment = align_center
        ws_arca.cell(row=r_arca_idx, column=4, value=item["comp_arca"]).alignment = align_center
        ws_arca.cell(row=r_arca_idx, column=5, value=item["cuit_arca"]).alignment = align_center
        ws_arca.cell(row=r_arca_idx, column=6, value=item["razon_arca"]).alignment = align_left

        c_m = ws_arca.cell(row=r_arca_idx, column=7, value=item["monto_arca"])
        c_m.number_format, c_m.alignment = fmt_currency, align_right

        for c in range(1, 8):
            ws_arca.cell(row=r_arca_idx, column=c).font = font_regular
            ws_arca.cell(row=r_arca_idx, column=c).border = thin_border
        r_arca_idx += 1

    ws_arca.cell(row=r_arca_idx, column=1, value=f"TOTAL PENDIENTE {ORI}").font = font_bold
    c_totarca = ws_arca.cell(row=r_arca_idx, column=7, value=f"=SUM(G4:G{r_arca_idx-1})")
    c_totarca.font, c_totarca.number_format, c_totarca.alignment = font_bold, fmt_currency, align_right

    for c in range(1, 8):
        ws_arca.cell(row=r_arca_idx, column=c).fill = fill_total
        ws_arca.cell(row=r_arca_idx, column=c).border = total_border

    row_arca_total = r_arca_idx
    cnt_arca = len(pendientes_arca)

    ws_mayor.cell(row=1, column=1, value=f"REGISTROS PENDIENTES EN MAYOR ({erp_seleccionado.upper()}){sufijo_titulo}").font = font_title
    headers_mayor = [
        "Nro Asiento",
        "Fecha Contable",
        "Detalle / Concepto",
        "Referencia / NroComp",
        "CUIT Cliente/Prov",
        "Razón Social / Cuenta",
        "Monto Registrado ($)",
    ]
    for col_idx, text in enumerate(headers_mayor, 1):
        c = ws_mayor.cell(row=3, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_header, fill_mayor_hdr, align_center

    r_mayor_idx = 4
    for item in pendientes_mayor:
        ws_mayor.cell(row=r_mayor_idx, column=1, value=item["asiento_mayor"]).alignment = align_center
        ws_mayor.cell(row=r_mayor_idx, column=2, value=item["fecha_mayor"]).alignment = align_center
        ws_mayor.cell(row=r_mayor_idx, column=3, value=item["detalle_mayor"]).alignment = align_left
        ws_mayor.cell(row=r_mayor_idx, column=4, value=item["ref_mayor"]).alignment = align_center
        ws_mayor.cell(row=r_mayor_idx, column=5, value=item["cuit_mayor"]).alignment = align_center
        ws_mayor.cell(row=r_mayor_idx, column=6, value=item["razon_mayor"]).alignment = align_left

        c_m = ws_mayor.cell(row=r_mayor_idx, column=7, value=item["monto_mayor"])
        c_m.number_format, c_m.alignment = fmt_currency, align_right

        for c in range(1, 8):
            ws_mayor.cell(row=r_mayor_idx, column=c).font = font_regular
            ws_mayor.cell(row=r_mayor_idx, column=c).border = thin_border
        r_mayor_idx += 1

    ws_mayor.cell(row=r_mayor_idx, column=1, value="TOTAL PENDIENTE MAYOR").font = font_bold
    c_totmayor = ws_mayor.cell(row=r_mayor_idx, column=7, value=f"=SUM(G4:G{r_mayor_idx-1})")
    c_totmayor.font, c_totmayor.number_format, c_totmayor.alignment = font_bold, fmt_currency, align_right

    for c in range(1, 8):
        ws_mayor.cell(row=r_mayor_idx, column=c).fill = fill_total
        ws_mayor.cell(row=r_mayor_idx, column=c).border = total_border

    row_mayor_total = r_mayor_idx
    cnt_mayor = len(pendientes_mayor)

    ws_resumen.cell(row=1, column=1, value=f"CONCILIACIÓN DE {concepto} - MAYOR {erp_seleccionado.upper()} VS {ORI}{sufijo_titulo}").font = font_title
    headers_resumen = [
        "Categoría / Pestaña",
        "Cantidad de Reg.",
        f"Total {ORI} ($)",
        "Total Mayor ($)",
        "Diferencia ($)",
    ]
    for col_idx, text in enumerate(headers_resumen, 1):
        c = ws_resumen.cell(row=4, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_header, fill_resumen_hdr, align_center

    resumen_rows = [
        (
            "1. Conciliaciones Exactas (1 a 1)",
            cnt_1a1,
            f"='Conciliadas 1a1'!F{row_1a1_total}",
            f"='Conciliadas 1a1'!L{row_1a1_total}",
            "=C5-D5",
        ),
        (
            "2. Conciliaciones por Lote / CUIT",
            cnt_lote,
            f"='Conciliadas por Lote'!E{row_lote_total}",
            f"='Conciliadas por Lote'!H{row_lote_total}",
            "=C6-D6",
        ),
        (
            f"3. Pendientes en {ORI} (Sin Mayor)",
            cnt_arca,
            f"={sheet_pend_arca}!G{row_arca_total}",
            0,
            "=C7-D7",
        ),
        (
            f"4. Pendientes en MAYOR (Sin {ORI})",
            cnt_mayor,
            0,
            f"={sheet_pend_mayor}!G{row_mayor_total}",
            "=C8-D8",
        ),
    ]

    for idx, rdata in enumerate(resumen_rows, 5):
        ws_resumen.cell(row=idx, column=1, value=rdata[0]).alignment = align_left
        ws_resumen.cell(row=idx, column=2, value=rdata[1]).alignment = align_center

        c_arca = ws_resumen.cell(row=idx, column=3, value=rdata[2])
        c_arca.number_format, c_arca.alignment = fmt_currency, align_right

        c_myr = ws_resumen.cell(row=idx, column=4, value=rdata[3])
        c_myr.number_format, c_myr.alignment = fmt_currency, align_right

        c_diff = ws_resumen.cell(row=idx, column=5, value=rdata[4])
        c_diff.number_format, c_diff.alignment = fmt_currency, align_right

        for c in range(1, 6):
            ws_resumen.cell(row=idx, column=c).font = font_regular
            ws_resumen.cell(row=idx, column=c).border = thin_border

    # TOTAL GENERAL
    ws_resumen.cell(row=9, column=1, value="TOTAL GENERAL EMANADO").font = font_bold
    ws_resumen.cell(row=9, column=2, value="=SUM(B5:B8)").font = font_bold
    ws_resumen.cell(row=9, column=2).alignment = align_center

    c_tot_arca = ws_resumen.cell(row=9, column=3, value="=SUM(C5:C8)")
    c_tot_arca.font, c_tot_arca.number_format, c_tot_arca.alignment = font_bold, fmt_currency, align_right

    c_tot_myr = ws_resumen.cell(row=9, column=4, value="=SUM(D5:D8)")
    c_tot_myr.font, c_tot_myr.number_format, c_tot_myr.alignment = font_bold, fmt_currency, align_right

    c_tot_diff = ws_resumen.cell(row=9, column=5, value="=C9-D9")
    c_tot_diff.font, c_tot_diff.number_format, c_tot_diff.alignment = font_bold, fmt_currency, align_right

    for c in range(1, 6):
        ws_resumen.cell(row=9, column=c).fill = fill_total
        ws_resumen.cell(row=9, column=c).border = total_border

    # Ajuste automático del ancho de columnas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]:
                    continue
                val_str = str(cell.value or "")
                if cell.number_format and "$" in cell.number_format:
                    val_str += "    "
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # Retornar tanto el Excel como los resúmenes para mostrar en Streamlit
    stats = {
        "cnt_1a1": cnt_1a1,
        "cnt_lote": cnt_lote,
        "cnt_arca": cnt_arca,
        "cnt_mayor": cnt_mayor,
        "conciliadas_1a1": conciliadas_1a1,
        "conciliadas_lote": conciliadas_lote,
        "pendientes_arca": pendientes_arca,
        "pendientes_mayor": pendientes_mayor,
    }

    return output, stats


def parse_autodealer_xls(file_bytes):
    """
    Interpreta un archivo .xls (que puede ser un reporte de texto disfrazado de Excel,
    un HTML disfrazado, o un Excel binario real) extrayendo sus datos línea por línea o fila por fila.

    El tipo de archivo se detecta por su contenido real (firma de bytes), no por la extensión.
    """
    file_bytes.seek(0)
    raw = file_bytes.read()
    file_bytes.seek(0)

    es_excel_binario = raw[:4] == b"\xd0\xcf\x11\xe0" or raw[:2] == b"PK"   # .xls (OLE2) o .xlsx (zip)
    inicio = raw[:2048].lower()
    es_html = (not es_excel_binario) and (b"<table" in inicio or b"<html" in inicio)

    lines = []

    if es_excel_binario or es_html:
        if es_html:
            df_raw = pd.read_html(io.StringIO(decodificar_texto(raw)))[0]
        else:
            try:
                # xlrd lee el formato viejo .xls
                df_raw = pd.read_excel(io.BytesIO(raw), header=None, dtype=str, engine="xlrd")
            except Exception:
                # Si no era .xls, probamos el motor por defecto (.xlsx)
                df_raw = pd.read_excel(io.BytesIO(raw), header=None, dtype=str)

        # Reconstruimos las filas del dataframe a texto para usar el motor de Regex
        for _, row in df_raw.iterrows():
            row_vals = []
            for val in row.values:
                if pd.notna(val) and str(val).strip() and str(val).lower() != 'nan':
                    v_str = str(val).strip()
                    # Normalizar fechas generadas por pandas de ser necesario
                    if re.match(r'^\d{4}-\d{2}-\d{2}( \d{2}:\d{2}:\d{2})?$', v_str):
                        try:
                            v_str = pd.to_datetime(v_str).strftime('%d/%m/%Y')
                        except Exception:
                            pass
                    row_vals.append(v_str)
            if row_vals:
                lines.append(" ".join(row_vals))
    else:
        # Reporte de texto plano (muy común en ERPs viejos que exportan a .xls / .prn)
        lines = decodificar_texto(raw).splitlines()

    data = []

    for line in lines:
        line = line.strip()
        # Patrón inicial: Validamos que la línea empiece con un Asiento y una Fecha sí o sí
        if not re.search(r'^\d+\s+\d{2}/\d{2}/\d{4}', line):
            continue

        # 1. Extraemos Asiento y Fecha
        asiento = re.search(r'^(\d+)', line).group(1)
        fecha = re.search(r'(\d{2}/\d{2}/\d{4})', line).group(1)

        # Cortamos lo que ya extrajimos del texto original
        resto = line.replace(asiento, "", 1).replace(fecha, "", 1).strip()

        # 2. Buscamos Referencia (Ej: PV 12132/0, US 12187/1)
        referencia = ""
        m_ref = re.search(r'^[A-Z]{2}\s\d+/\d', resto)
        if m_ref:
            referencia = m_ref.group(0)
            resto = resto.replace(referencia, "", 1).strip()

        # 3. Buscamos Detalle / Recibo (Ej: RC-X-0100-00104458)
        detalle = ""
        m_det = re.search(r'^[A-Za-z0-9]+\-[A-Za-z0-9\-]+', resto)
        if m_det:
            detalle = m_det.group(0)
            resto = resto.replace(detalle, "", 1).strip()

        # 4. Extraemos todos los montos con decimales desde el texto restante
        # Usamos un patrón flexible para soportar formatos con comas, puntos o reconstruidos por pandas
        amt_pattern = r'\b(?:\d{1,3}(?:[.,]\d{3})+|\d+)[.,]\d{1,2}\b'
        montos = re.findall(amt_pattern, resto)

        # Quitamos los montos del texto para aislar la Entidad y el Centro de Costo
        for m in montos:
            resto = resto.replace(m, "", 1)

        resto = resto.strip()

        # 5. Separamos Centro de Costo (última palabra pegada al final) de la Entidad
        centro_costo = ""
        entidad = resto
        m_cc = re.search(r'\s+([A-Za-z]+)$', resto)
        if m_cc:
            centro_costo = m_cc.group(1)
            entidad = resto[:m_cc.start()].strip()

        # Limpiamos espacios dobles accidentales en la entidad resultante
        entidad = re.sub(r'\s+', ' ', entidad).strip()

        # Formateamos REF corta (quitando lo que está después de la barra)
        ref_corta = referencia.split('/')[0] if referencia else ""

        # 6. Acomodamos DEBE, HABER, SALDO según la cantidad de montos extraídos
        debe, haber, saldo = "", "", ""
        if len(montos) >= 3:
            debe, haber, saldo = montos[-3], montos[-2], montos[-1]
        elif len(montos) == 2:
            # Comportamiento común del PRN: asume DEBE y SALDO
            debe, saldo = montos[0], montos[1]
        elif len(montos) == 1:
            saldo = montos[0]

        data.append({
            "ASIENTO": asiento,
            "FECHA": fecha,
            "Referencia": referencia,
            "DETALLE": detalle,
            "REF": ref_corta,
            "ENTIDAD": entidad,
            "DEBE": debe,
            "HABER": haber,
            "SALDO": saldo,
            "CENTRO DE COSTO": centro_costo
        })

    df = pd.DataFrame(data)

    # Guardar en memoria como Excel
    output = io.BytesIO()
    df.to_excel(output, index=False, sheet_name="Mayor Formateado")
    output.seek(0)
    return output


def main():
    st.sidebar.title("🛠️ Herramientas")
    modo_app = st.sidebar.radio(
        "Selecciona un módulo:",
        ["🔄 Conciliador de Retenciones", "🧾 Conciliador de Percepciones IIBB", "✨ Formatear Mayor (XLS)"]
    )
    st.sidebar.markdown("---")

    if modo_app == "🔄 Conciliador de Retenciones":
        st.title("📊 Conciliador Contable de Retenciones (ARCA vs ERP)")
        st.markdown(
            "Selecciona el **ERP de origen**, carga los reportes de **Mayor** y **Mis Retenciones (ARCA)**, "
            "y genera el cruce automático con validación por **CUIT y Razón Social**."
        )

        st.warning("⚠️ **ATENCIÓN:** Solo se admiten archivos en formato moderno **.xlsx**. Si tus archivos son `.xls` (muy común en ARCA/AFIP), ábrelos en Excel y guárdalos como `.xlsx` antes de subirlos.", icon="⚠️")

        with st.sidebar:
            st.header("⚙️ Configuración del Cruce")
            erp_seleccionado = st.radio(
                "Selecciona el ERP del Mayor:",
                options=["AutoDealer", "Bejerman"],
                help="Adapta dinámicamente las referencias de los encabezados según el sistema contable emisor."
            )

            st.markdown("---")
            st.markdown("### 📌 Mapeo activo")
            cfg = ERP_CONFIGS[erp_seleccionado]
            st.caption(f"**Asiento:** `{cfg.get('asiento')}`")
            st.caption(f"**Fecha:** `{cfg.get('fecha')}`")
            st.caption(f"**Detalle / Concepto:** `{cfg.get('detalle')}`")
            st.caption(f"**Referencia:** `{cfg.get('referencia')}`")
            if cfg.get("ref_secundaria"):
                st.caption(f"**Ref Secundario:** `{cfg.get('ref_secundaria')}`")
            st.caption(f"**Razón Social:** `{cfg.get('entidad')}`")
            st.caption(f"**CUIT:** `{cfg.get('cuit')}`")
            st.caption(f"**Debe:** `{cfg.get('debe')}` | **Haber:** `{cfg.get('haber')}`")

        col1, col2 = st.columns(2)

        with col1:
            file_mayor = st.file_uploader(
                f"Subir Mayor (.xlsx) ({erp_seleccionado})", type=["xlsx"], key="mayor"
            )

        with col2:
            file_arca = st.file_uploader(
                "Subir MisRetenciones (.xlsx) (ARCA)", type=["xlsx"], key="arca"
            )

        if file_mayor and file_arca:
            st.info(f"💡 Listo para realizar el cruce utilizando la estructura de **{erp_seleccionado}**.")

            if st.button("⚙️ Procesar Conciliación", type="primary", use_container_width=True):
                with st.spinner("Procesando retenciones y aplicando cruce por CUIT..."):
                    try:
                        excel_bytes, stats = procesar_archivos(file_mayor, file_arca, erp_seleccionado)

                        st.success("✅ ¡Conciliación completada con éxito!")

                        m1, m2, m3, m4 = st.columns(4)
                        m1.metric("Conciliados 1 a 1", stats["cnt_1a1"])
                        m2.metric("Conciliados por Lote", stats["cnt_lote"])
                        m3.metric("Pendientes ARCA", stats["cnt_arca"])
                        m4.metric("Pendientes Mayor", stats["cnt_mayor"])

                        st.download_button(
                            label="📥 Descargar Reporte Conciliado en Excel",
                            data=excel_bytes,
                            file_name=f"Conciliacion_Retenciones_{erp_seleccionado}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                        )

                        st.markdown("### 🔍 Previsualización de Resultados")
                        tab1, tab2, tab3, tab4 = st.tabs([
                            "Exactas 1a1", "Por Lote / CUIT", "Pendientes ARCA", "Pendientes Mayor"
                        ])

                        with tab1:
                            if stats["conciliadas_1a1"]:
                                st.dataframe(pd.DataFrame(stats["conciliadas_1a1"]), use_container_width=True)
                            else:
                                st.info("No se encontraron registros 1 a 1 exactos.")

                        with tab2:
                            if stats["conciliadas_lote"]:
                                st.dataframe(pd.DataFrame(stats["conciliadas_lote"]), use_container_width=True)
                            else:
                                st.info("No se encontraron registros acumulados por lote.")

                        with tab3:
                            if stats["pendientes_arca"]:
                                st.dataframe(pd.DataFrame(stats["pendientes_arca"]), use_container_width=True)
                            else:
                                st.success("🎉 ¡No hay retenciones pendientes en ARCA!")

                        with tab4:
                            if stats["pendientes_mayor"]:
                                st.dataframe(pd.DataFrame(stats["pendientes_mayor"]), use_container_width=True)
                            else:
                                st.success("🎉 ¡No hay registros pendientes en el Mayor!")
                    except ValueError as ve:
                        st.error(f"⚠️ {str(ve)}", icon="🛑")
                    except Exception as error:
                        st.error(f"❌ Ocurrió un error inesperado al procesar los archivos: {str(error)}", icon="❌")

    elif modo_app == "🧾 Conciliador de Percepciones IIBB":
        st.title("🧾 Conciliador de Percepciones IIBB (Listado del sistema vs Mayor Bejerman)")
        st.markdown(
            "Carga el **Mayor de Bejerman** (una sola jurisdicción por reporte), el **.txt de percepciones** del mes "
            "y el **listado de CUIT Proveedores**. Luego elegí la jurisdicción a conciliar."
        )
        st.warning("⚠️ El Mayor debe estar en formato **.xlsx**. El cruce usa la columna **ImpTotal_Loc** (las NC van en negativo).", icon="⚠️")

        with st.sidebar:
            st.header("⚙️ Configuración del Cruce")
            st.caption("ERP del Mayor: **Bejerman** (fijo en este módulo)")
            st.markdown("---")
            st.markdown("### 📌 Mapeo activo")
            cfg = ERP_CONFIGS["Bejerman"]
            st.caption(f"**Asiento:** `{cfg.get('asiento')}`")
            st.caption(f"**Fecha:** `{cfg.get('fecha')}`")
            st.caption(f"**Detalle / Concepto:** `{cfg.get('detalle')}`")
            st.caption(f"**Ref Secundario:** `{cfg.get('ref_secundaria')}`")
            st.caption(f"**Razón Social:** `{cfg.get('entidad')}`")
            st.caption(f"**CUIT:** `{cfg.get('cuit')}`")
            st.caption(f"**Monto de cruce:** `{cfg.get('saldo')}`")

        c1, c2, c3 = st.columns(3)
        with c1:
            file_mayor_p = st.file_uploader("Subir Mayor Bejerman (.xlsx)", type=["xlsx"], key="mayor_iibb")
        with c2:
            file_txt = st.file_uploader("Subir listado de percepciones (.txt)", type=["txt"], key="txt_iibb")
        with c3:
            file_prov = st.file_uploader("Subir CUIT Proveedores (.txt)", type=["txt"], key="prov_iibb")

        df_txt = None
        jurisdiccion = None

        if file_txt:
            try:
                df_txt = parse_percepciones_txt(file_txt.getvalue())
            except Exception as e:
                st.error(f"❌ No se pudo interpretar el .txt de percepciones: {e}")

            if df_txt is not None and df_txt.empty:
                st.error("❌ No se encontraron líneas con el formato esperado en el .txt de percepciones.")
                df_txt = None

        if df_txt is not None:
            resumen_jur = df_txt.groupby("Jurisdicción")["Monto"].agg(["count", "sum"])
            jurisdicciones = sorted(resumen_jur.index.tolist())

            def _fmt_jur(j):
                n = int(resumen_jur.loc[j, "count"])
                total = resumen_jur.loc[j, "sum"]
                return f"{j}  —  {n} registros  —  $ {total:,.2f}"

            jurisdiccion = st.selectbox(
                "¿Qué jurisdicción necesitás conciliar?",
                options=jurisdicciones,
                format_func=_fmt_jur,
            )

            with st.expander("👁️ Ver el .txt ya formateado (todas las jurisdicciones)"):
                st.dataframe(df_txt, use_container_width=True)
                buf = io.BytesIO()
                df_txt.to_excel(buf, index=False, sheet_name="Percepciones")
                buf.seek(0)
                st.download_button(
                    "📥 Descargar .txt formateado (.xlsx)",
                    data=buf,
                    file_name="Percepciones_Formateadas.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

        if file_mayor_p and df_txt is not None and file_prov and jurisdiccion:
            st.info(f"💡 Listo para conciliar la jurisdicción **{jurisdiccion}** contra el Mayor de Bejerman.")

            if st.button("⚙️ Procesar Conciliación", type="primary", use_container_width=True):
                with st.spinner("Procesando percepciones y aplicando cruce..."):
                    try:
                        padron = cargar_padron_proveedores(file_prov.getvalue())
                        df_j = df_txt[df_txt["Jurisdicción"] == jurisdiccion].copy()
                        df_j = agregar_razon_social(df_j, padron)
                        df_motor = construir_df_motor(df_j)

                        excel_bytes, stats = procesar_archivos(
                            file_mayor_p,
                            None,
                            "Bejerman",
                            df_arca_override=df_motor,
                            origen_label="IIBB",
                            concepto="PERCEPCIONES IIBB",
                            subtitulo=f"JURISDICCIÓN {jurisdiccion}",
                            priorizar_total_mayor=True,
                        )

                        st.success(f"✅ ¡Conciliación de la jurisdicción {jurisdiccion} completada!")

                        sin_rs = df_j[df_j["Razón Social"] == SIN_RAZON_SOCIAL]["CUIT"].unique().tolist()
                        if sin_rs:
                            st.warning(
                                f"⚠️ {len(sin_rs)} CUIT del .txt no figuran en el listado de proveedores "
                                f"(se informan como '{SIN_RAZON_SOCIAL}'): " + ", ".join(sin_rs)
                            )

                        m1, m2, m3, m4 = st.columns(4)
                        m1.metric("Conciliados 1 a 1", stats["cnt_1a1"])
                        m2.metric("Conciliados por Lote", stats["cnt_lote"])
                        m3.metric("Pendientes IIBB", stats["cnt_arca"])
                        m4.metric("Pendientes Mayor", stats["cnt_mayor"])

                        st.download_button(
                            label="📥 Descargar Reporte Conciliado en Excel",
                            data=excel_bytes,
                            file_name=f"Conciliacion_Percepciones_IIBB_{jurisdiccion}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                        )

                        st.markdown("### 🔍 Previsualización de Resultados")
                        tab1, tab2, tab3, tab4 = st.tabs([
                            "Exactas 1a1", "Por Lote / CUIT", "Pendientes IIBB", "Pendientes Mayor"
                        ])
                        with tab1:
                            if stats["conciliadas_1a1"]:
                                st.dataframe(pd.DataFrame(stats["conciliadas_1a1"]), use_container_width=True)
                            else:
                                st.info("No se encontraron registros 1 a 1 exactos.")
                        with tab2:
                            if stats["conciliadas_lote"]:
                                st.dataframe(pd.DataFrame(stats["conciliadas_lote"]), use_container_width=True)
                            else:
                                st.info("No se encontraron registros acumulados por lote.")
                        with tab3:
                            if stats["pendientes_arca"]:
                                st.dataframe(pd.DataFrame(stats["pendientes_arca"]), use_container_width=True)
                            else:
                                st.success("🎉 ¡No hay percepciones pendientes en el listado!")
                        with tab4:
                            if stats["pendientes_mayor"]:
                                st.dataframe(pd.DataFrame(stats["pendientes_mayor"]), use_container_width=True)
                            else:
                                st.success("🎉 ¡No hay registros pendientes en el Mayor!")
                    except ValueError as ve:
                        st.error(f"⚠️ {str(ve)}", icon="🛑")
                    except Exception as error:
                        st.error(f"❌ Ocurrió un error inesperado al procesar los archivos: {str(error)}", icon="❌")

    elif modo_app == "✨ Formatear Mayor (XLS)":
        st.title("✨ Formateador Automático de Mayor (AutoDealer)")
        st.markdown(
            "Sube tu archivo crudo **.xls** exportado del sistema. "
            "El motor inteligente extraerá números de comprobante, entidades sociales de longitud variable y montos, "
            "generando un **Excel tabulado (.xlsx)** listo para el conciliador."
        )

        file_xls = st.file_uploader("📂 Sube el archivo crudo del Mayor", type=["xls", "xlsx", "txt", "prn"])

        if file_xls:
            st.info("💡 Archivo leído. Listo para ser interpretado y estructurado.")
            if st.button("🪄 Ejecutar Formateo", type="primary", use_container_width=True):
                with st.spinner("Desglosando registros del mayor y limpiando columnas..."):
                    try:
                        excel_formateado = parse_autodealer_xls(file_xls)
                        st.success("✅ ¡Archivo procesado y estructurado con éxito!")

                        st.download_button(
                            label="📥 Descargar Mayor Formateado (.xlsx)",
                            data=excel_formateado,
                            file_name="Mayor_AutoDealer_Formateado.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True
                        )

                        st.markdown("⚠️ *Recomendación: Verifica que el formato descargado sea el esperado (especialmente la separación DEBE/HABER/SALDO) antes de importarlo al conciliador.*")
                    except Exception as e:
                        st.error(f"❌ Ocurrió un error al procesar el archivo: {str(e)}")


if __name__ == "__main__":
    main()
