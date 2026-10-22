"""Excel del Mayor formateado: encabezado, formatos numéricos, filtros y anchos."""
import io

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from reports.styles import ALIGN_CENTER, FILL_RESUMEN, FONT_HEADER

_COLUMNAS_MONTO = ("DEBE", "HABER", "SALDO")
_ANCHO_MIN, _ANCHO_MAX = 12, 60


def generar_excel_mayor(df, nombre_hoja="Mayor Formateado"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill, c.alignment = FONT_HEADER, FILL_RESUMEN, ALIGN_CENTER

    filas = df.astype(object).where(df.notna(), None).values.tolist()
    for fila in filas:
        ws.append(fila)

    for nombre in _COLUMNAS_MONTO:
        if nombre in df.columns:
            col = df.columns.get_loc(nombre) + 1
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0.00"

    for idx, nombre in enumerate(df.columns, 1):
        largo = max([len(str(nombre))] + [len(str(v)) for v in df[nombre].dropna().head(500)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(largo + 3, _ANCHO_MIN), _ANCHO_MAX)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
