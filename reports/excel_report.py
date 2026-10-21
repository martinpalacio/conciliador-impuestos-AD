"""Genera el Excel de conciliación: una función por hoja + un orquestador."""
import io

import openpyxl

from reports.excel_helpers import (
    autoajustar_columnas, escribir_encabezados, escribir_fila_total, escribir_filas,
    escribir_titulo, formula_suma,
)
from reports.styles import (
    ALIGN_CENTER as C, ALIGN_LEFT as L, ALIGN_RIGHT as R, FILL_1A1, FILL_LOTE,
    FILL_PEND_EXTERNO, FILL_PEND_MAYOR, FILL_RESUMEN, FILL_TOTAL, FMT_CURRENCY,
    FONT_BOLD, FONT_REGULAR, THIN_BORDER, TOTAL_BORDER,
)

HOJA_RESUMEN = "Resumen General"
HOJA_1A1 = "Conciliadas 1a1"
HOJA_LOTE = "Conciliadas por Lote"
HOJA_PEND_MAYOR = "Pendientes_MAYOR"


def _hoja_1a1(ws, items, ori, erp, sufijo):
    escribir_titulo(ws, f"PARTIDAS CONCILIADAS EXACTAS 1 A 1 ({erp.upper()}){sufijo}")
    lbl_cert = "Certif." if ori == "ARCA" else "Nro Comp."
    escribir_encabezados(ws, 3, [
        f"{lbl_cert} {ori}", f"Fecha {ori}", f"Comprobante {ori}", f"CUIT {ori}",
        f"Razón Social {ori}", f"Monto {ori} ($)", "Asiento Mayor", "Fecha Mayor",
        "Detalle / Comp. Mayor", "CUIT Mayor", "Razón Social Mayor", "Monto Mayor ($)",
        "Diferencia ($)",
    ], FILL_1A1)

    columnas = [
        ("cert_arca", C, False), ("fecha_arca", C, False), ("comp_arca", C, False),
        ("cuit_arca", C, False), ("razon_arca", L, False), ("monto_arca", R, True),
        ("asiento_mayor", C, False), ("fecha_mayor", C, False), ("comp_mayor", L, False),
        ("cuit_mayor", C, False), ("razon_mayor", L, False), ("monto_mayor", R, True),
        (lambda f: f"=F{f}-L{f}", R, True),
    ]
    fila_total = escribir_filas(ws, items, columnas, 4)
    escribir_fila_total(
        ws, fila_total, 13, "TOTAL CONCILIADO 1 A 1",
        {6: formula_suma("F", 4, fila_total), 12: formula_suma("L", 4, fila_total),
         13: formula_suma("M", 4, fila_total)},
        etiqueta_a_izquierda=True, alinear_derecha=False,
    )
    return fila_total


def _hoja_lote(ws, items, ori, erp, sufijo):
    escribir_titulo(ws, f"PARTIDAS CONCILIADAS POR LOTE / CUIT ({erp.upper()}){sufijo}")
    escribir_encabezados(ws, 4, [
        "Razón Social / Empresa", "CUIT",
        f"{'Certificados' if ori == 'ARCA' else 'Comprobantes'} {ori} Incluidos",
        f"Rango Fecha {ori}", f"Monto Total {ori} ($)", "Asientos Mayor Incluidos",
        "Rango Fecha Mayor", "Monto Total Mayor ($)", "Diferencia ($)", "Observaciones",
    ], FILL_LOTE)

    columnas = [
        ("empresa", L, False), ("cuit", C, False), ("certs_arca", C, False),
        ("fecha_arca", C, False), ("monto_arca", R, True), ("asientos_mayor", C, False),
        ("fecha_mayor", C, False), ("monto_mayor", R, True),
        (lambda f: f"=E{f}-H{f}", R, True), ("obs", L, False),
    ]
    fila_total = escribir_filas(ws, items, columnas, 5)
    escribir_fila_total(
        ws, fila_total, 10, "TOTAL CONCILIADO POR LOTE",
        {5: formula_suma("E", 5, fila_total), 8: formula_suma("H", 5, fila_total),
         9: formula_suma("I", 5, fila_total)},
    )
    return fila_total


def _hoja_pendientes_externo(ws, items, ori, concepto, sufijo):
    escribir_titulo(ws, f"{concepto} PENDIENTES EN {ori} (NO REGISTRADAS EN MAYOR){sufijo}")
    escribir_encabezados(ws, 3, [
        "Nro Certificado" if ori == "ARCA" else "Nro Comprobante",
        "Fecha Reg.", "Tipo Comprobante", "Nro Comprobante", "CUIT Agente",
        "Razón Social Agente", "Monto Retención ($)" if ori == "ARCA" else "Monto ($)",
    ], FILL_PEND_EXTERNO)

    columnas = [
        ("cert_arca", C, False), ("fecha_arca", C, False), ("tipo_comp", C, False),
        ("comp_arca", C, False), ("cuit_arca", C, False), ("razon_arca", L, False),
        ("monto_arca", R, True),
    ]
    fila_total = escribir_filas(ws, items, columnas, 4)
    escribir_fila_total(ws, fila_total, 7, f"TOTAL PENDIENTE {ori}",
                        {7: formula_suma("G", 4, fila_total)})
    return fila_total


def _hoja_pendientes_mayor(ws, items, erp, sufijo):
    escribir_titulo(ws, f"REGISTROS PENDIENTES EN MAYOR ({erp.upper()}){sufijo}")
    escribir_encabezados(ws, 3, [
        "Nro Asiento", "Fecha Contable", "Detalle / Concepto", "Referencia / NroComp",
        "CUIT Cliente/Prov", "Razón Social / Cuenta", "Monto Registrado ($)",
    ], FILL_PEND_MAYOR)

    columnas = [
        ("asiento_mayor", C, False), ("fecha_mayor", C, False), ("detalle_mayor", L, False),
        ("ref_mayor", C, False), ("cuit_mayor", C, False), ("razon_mayor", L, False),
        ("monto_mayor", R, True),
    ]
    fila_total = escribir_filas(ws, items, columnas, 4)
    escribir_fila_total(ws, fila_total, 7, "TOTAL PENDIENTE MAYOR",
                        {7: formula_suma("G", 4, fila_total)})
    return fila_total


def _hoja_resumen(ws, ori, erp, concepto, sufijo, filas):
    """`filas`: dict con cantidades (cnt_*) y filas de total (row_*) de las otras hojas."""
    hoja_pend_ext = f"Pendientes_{ori}"
    escribir_titulo(ws, f"CONCILIACIÓN DE {concepto} - MAYOR {erp.upper()} VS {ori}{sufijo}")
    escribir_encabezados(ws, 4, [
        "Categoría / Pestaña", "Cantidad de Reg.", f"Total {ori} ($)", "Total Mayor ($)",
        "Diferencia ($)",
    ], FILL_RESUMEN)

    resumen_rows = [
        ("1. Conciliaciones Exactas (1 a 1)", filas["cnt_1a1"],
         f"='{HOJA_1A1}'!F{filas['row_1a1']}", f"='{HOJA_1A1}'!L{filas['row_1a1']}", "=C5-D5"),
        ("2. Conciliaciones por Lote / CUIT", filas["cnt_lote"],
         f"='{HOJA_LOTE}'!E{filas['row_lote']}", f"='{HOJA_LOTE}'!H{filas['row_lote']}", "=C6-D6"),
        (f"3. Pendientes en {ori} (Sin Mayor)", filas["cnt_arca"],
         f"={hoja_pend_ext}!G{filas['row_arca']}", 0, "=C7-D7"),
        (f"4. Pendientes en MAYOR (Sin {ori})", filas["cnt_mayor"],
         0, f"={HOJA_PEND_MAYOR}!G{filas['row_mayor']}", "=C8-D8"),
    ]

    for idx, (label, cnt, val_ext, val_myr, diff) in enumerate(resumen_rows, 5):
        ws.cell(row=idx, column=1, value=label).alignment = L
        ws.cell(row=idx, column=2, value=cnt).alignment = C
        for col, valor in ((3, val_ext), (4, val_myr), (5, diff)):
            c = ws.cell(row=idx, column=col, value=valor)
            c.number_format, c.alignment = FMT_CURRENCY, R
        for c in range(1, 6):
            ws.cell(row=idx, column=c).font = FONT_REGULAR
            ws.cell(row=idx, column=c).border = THIN_BORDER

    ws.cell(row=9, column=1, value="TOTAL GENERAL EMANADO").font = FONT_BOLD
    ws.cell(row=9, column=2, value="=SUM(B5:B8)").font = FONT_BOLD
    ws.cell(row=9, column=2).alignment = C
    for col, formula in ((3, "=SUM(C5:C8)"), (4, "=SUM(D5:D8)"), (5, "=C9-D9")):
        c = ws.cell(row=9, column=col, value=formula)
        c.font, c.number_format, c.alignment = FONT_BOLD, FMT_CURRENCY, R
    for c in range(1, 6):
        ws.cell(row=9, column=c).fill = FILL_TOTAL
        ws.cell(row=9, column=c).border = TOTAL_BORDER


def generar_excel(resultado, erp, origen_label, concepto, subtitulo=""):
    """
    Arma el libro completo a partir del resultado del motor (dict con las 4 listas).
    Devuelve un BytesIO listo para descargar.
    """
    ori = origen_label
    sufijo = f" - {subtitulo}" if subtitulo else ""
    hoja_pend_ext = f"Pendientes_{ori}"

    wb = openpyxl.Workbook()
    ws_resumen = wb.active
    ws_resumen.title = HOJA_RESUMEN
    ws_1a1 = wb.create_sheet(title=HOJA_1A1)
    ws_lote = wb.create_sheet(title=HOJA_LOTE)
    ws_ext = wb.create_sheet(title=hoja_pend_ext)
    ws_mayor = wb.create_sheet(title=HOJA_PEND_MAYOR)

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True

    filas = {
        "row_1a1": _hoja_1a1(ws_1a1, resultado["conciliadas_1a1"], ori, erp, sufijo),
        "row_lote": _hoja_lote(ws_lote, resultado["conciliadas_lote"], ori, erp, sufijo),
        "row_arca": _hoja_pendientes_externo(ws_ext, resultado["pendientes_arca"], ori, concepto, sufijo),
        "row_mayor": _hoja_pendientes_mayor(ws_mayor, resultado["pendientes_mayor"], erp, sufijo),
        "cnt_1a1": len(resultado["conciliadas_1a1"]),
        "cnt_lote": len(resultado["conciliadas_lote"]),
        "cnt_arca": len(resultado["pendientes_arca"]),
        "cnt_mayor": len(resultado["pendientes_mayor"]),
    }
    _hoja_resumen(ws_resumen, ori, erp, concepto, sufijo, filas)

    autoajustar_columnas(wb)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
