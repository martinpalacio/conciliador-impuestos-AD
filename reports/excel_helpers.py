"""Helpers genéricos para escribir tablas con estilo en una hoja de openpyxl."""
from openpyxl.utils import get_column_letter

from reports.styles import (
    ALIGN_LEFT, ALIGN_RIGHT, ALIGN_CENTER, FILL_TOTAL, FMT_CURRENCY, FONT_BOLD,
    FONT_HEADER, FONT_REGULAR, FONT_TITLE, THIN_BORDER, TOTAL_BORDER,
)


def escribir_titulo(ws, texto):
    ws.cell(row=1, column=1, value=texto).font = FONT_TITLE


def escribir_encabezados(ws, fila, headers, fill):
    for col_idx, text in enumerate(headers, 1):
        c = ws.cell(row=fila, column=col_idx, value=text)
        c.font, c.fill, c.alignment = FONT_HEADER, fill, ALIGN_CENTER


def escribir_filas(ws, items, columnas, fila_inicio):
    """
    Escribe una fila por item. `columnas` es una lista de (fuente, alineación, es_moneda):
    fuente es la clave del dict del item o una función(fila) -> valor (para fórmulas).
    Devuelve la fila siguiente a la última escrita (donde va el total).
    """
    fila = fila_inicio
    for item in items:
        for col_idx, (fuente, alineacion, es_moneda) in enumerate(columnas, 1):
            valor = fuente(fila) if callable(fuente) else item[fuente]
            c = ws.cell(row=fila, column=col_idx, value=valor)
            c.alignment = alineacion
            if es_moneda:
                c.number_format = FMT_CURRENCY
            c.font = FONT_REGULAR
            c.border = THIN_BORDER
        fila += 1
    return fila


def formula_suma(letra_col, fila_ini, fila_total):
    """
    =SUM(...) sobre el rango de datos. Si no hay filas devuelve 0: un rango vacío
    (p. ej. F4:F3) se invertiría a F3:F4 e incluiría la propia celda del total (referencia circular).
    """
    if fila_total <= fila_ini:
        return 0
    return f"=SUM({letra_col}{fila_ini}:{letra_col}{fila_total - 1})"


def escribir_fila_total(ws, fila, ncols, etiqueta, celdas, etiqueta_a_izquierda=False, alinear_derecha=True):
    """Fila de totales: etiqueta en col 1 y `celdas` = {nro_columna: valor_o_formula} en formato moneda."""
    c_label = ws.cell(row=fila, column=1, value=etiqueta)
    c_label.font = FONT_BOLD
    if etiqueta_a_izquierda:
        c_label.alignment = ALIGN_LEFT

    for col, valor in celdas.items():
        c = ws.cell(row=fila, column=col, value=valor)
        c.font, c.number_format = FONT_BOLD, FMT_CURRENCY
        if alinear_derecha:
            c.alignment = ALIGN_RIGHT

    for c in range(1, ncols + 1):
        ws.cell(row=fila, column=c).fill = FILL_TOTAL
        ws.cell(row=fila, column=c).border = TOTAL_BORDER


def autoajustar_columnas(wb):
    """Ajusta el ancho de cada columna según su contenido (ignora las filas de título)."""
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in (1, 2):
                    continue
                val_str = str(cell.value or "")
                if cell.number_format and "$" in cell.number_format:
                    val_str += "    "
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
